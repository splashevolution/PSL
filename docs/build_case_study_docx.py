#!/usr/bin/env python3
"""Build the PVM case-study DOCX from its Markdown source."""

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ACCENT = "173B57"
LIGHT_ACCENT = "E9F0F4"
TEXT = RGBColor(35, 44, 52)


def set_font(run, name="Aptos", size=None, bold=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:cs"), "Nirmala UI")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def set_style_font(style, name="Aptos", size=None, bold=None, color=None):
    style.font.name = name
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    style._element.get_or_add_rPr().rFonts.set(qn("w:cs"), "Nirmala UI")
    if size is not None:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if color is not None:
        style.font.color.rgb = color


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_cell_margins(cell):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for edge in ("top", "start", "bottom", "end"):
        tag = "w:" + edge
        node = margins.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            margins.append(node)
        node.set(qn("w:w"), "120")
        node.set(qn("w:type"), "dxa")


def parse_inline(paragraph, text, style=None):
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*|\*[^*]+\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_font(run, "Consolas", 9.5, color=RGBColor(28, 60, 81))
        elif part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_font(run, bold=True, color=TEXT)
        elif part.startswith("*") and part.endswith("*"):
            run = paragraph.add_run(part[1:-1])
            set_font(run, color=TEXT)
            run.italic = True
        else:
            run = paragraph.add_run(part)
            set_font(run, color=TEXT)
    if style:
        paragraph.style = style


def add_table(document, rows):
    table = document.add_table(rows=0, cols=len(rows[0]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row_index, row_values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(row_values):
            cell = cells[index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            parse_inline(paragraph, value)
            for run in paragraph.runs:
                set_font(run, size=9, bold=(row_index == 0),
                         color=(RGBColor(255, 255, 255) if row_index == 0 else TEXT))
            shade_cell(cell, ACCENT if row_index == 0 else ("F6F8F9" if row_index % 2 == 0 else "FFFFFF"))
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def build_document(markdown_path, output_path):
    content = Path(markdown_path).read_text(encoding="utf-8").splitlines()
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.84)
    section.right_margin = Inches(0.84)

    styles = document.styles
    set_style_font(styles["Normal"], size=10.5, color=TEXT)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Normal"].paragraph_format.line_spacing = 1.12
    for name, size in (("Title", 26), ("Heading 1", 18), ("Heading 2", 13.5), ("Heading 3", 11.5)):
        style = styles[name]
        set_style_font(style, size=size, bold=True, color=RGBColor(23, 59, 87))
        style.paragraph_format.space_before = Pt(14 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(7)

    header = section.header.paragraphs[0]
    header.text = "PSL: PĀṆINIAN SYSTEMS LANGUAGE  |  TECHNICAL CASE STUDY"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(header.runs[0], size=8.5, bold=True, color=RGBColor(100, 110, 118))

    footer = section.footer.paragraphs[0]
    footer.text = "Reproducible prototype evidence - 23 May 2026"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(footer.runs[0], size=8, color=RGBColor(100, 110, 118))

    index = 0
    code_mode = False
    code_lines = []
    paragraph_lines = []

    def flush_paragraph():
        if paragraph_lines:
            paragraph = document.add_paragraph()
            parse_inline(paragraph, " ".join(paragraph_lines))
            paragraph_lines.clear()

    while index < len(content):
        line = content[index]
        if line.startswith("```"):
            flush_paragraph()
            if code_mode:
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.left_indent = Inches(0.18)
                paragraph.paragraph_format.right_indent = Inches(0.18)
                paragraph.paragraph_format.space_before = Pt(4)
                paragraph.paragraph_format.space_after = Pt(8)
                for n, code_line in enumerate(code_lines):
                    run = paragraph.add_run(code_line + ("\n" if n < len(code_lines) - 1 else ""))
                    set_font(run, "Consolas", 8.7, color=RGBColor(35, 44, 52))
                p_pr = paragraph._p.get_or_add_pPr()
                shd = OxmlElement("w:shd")
                shd.set(qn("w:fill"), LIGHT_ACCENT)
                p_pr.append(shd)
                code_lines.clear()
                code_mode = False
            else:
                code_mode = True
            index += 1
            continue
        if code_mode:
            code_lines.append(line)
            index += 1
            continue
        if not line.strip():
            flush_paragraph()
            index += 1
            continue
        if line.startswith("| ") and index + 1 < len(content) and content[index + 1].startswith("| ---"):
            flush_paragraph()
            rows = []
            index += 2
            header_values = [value.strip() for value in line.strip("|").split("|")]
            rows.append(header_values)
            while index < len(content) and content[index].startswith("| "):
                rows.append([value.strip() for value in content[index].strip("|").split("|")])
                index += 1
            add_table(document, rows)
            continue
        if line.startswith("# "):
            flush_paragraph()
            paragraph = document.add_paragraph(style="Title")
            parse_inline(paragraph, line[2:])
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        elif line.startswith("## "):
            flush_paragraph()
            paragraph = document.add_paragraph(style="Heading 1")
            parse_inline(paragraph, line[3:])
        elif line.startswith("### "):
            flush_paragraph()
            paragraph = document.add_paragraph(style="Heading 2")
            parse_inline(paragraph, line[4:])
        elif line.startswith("> "):
            flush_paragraph()
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.28)
            paragraph.paragraph_format.right_indent = Inches(0.2)
            paragraph.paragraph_format.space_before = Pt(4)
            paragraph.paragraph_format.space_after = Pt(10)
            parse_inline(paragraph, line[2:])
            for run in paragraph.runs:
                run.italic = True
                set_font(run, size=11, color=RGBColor(23, 59, 87))
        elif re.match(r"^\d+\.\s", line):
            flush_paragraph()
            paragraph = document.add_paragraph(style="List Number")
            parse_inline(paragraph, re.sub(r"^\d+\.\s", "", line))
        elif line.startswith("- "):
            flush_paragraph()
            paragraph = document.add_paragraph(style="List Bullet")
            parse_inline(paragraph, line[2:])
        elif line.startswith("**") and line.endswith("**"):
            flush_paragraph()
            paragraph = document.add_paragraph()
            parse_inline(paragraph, line)
        else:
            paragraph_lines.append(line.strip())
        index += 1

    flush_paragraph()
    document.core_properties.title = "From Karaka Roles to Executable Isolation Semantics"
    document.core_properties.subject = "Paninian Virtual Machine RV32 proof case study"
    document.core_properties.author = "Praveen Kumar"
    document.save(output_path)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("Usage: build_case_study_docx.py <source.md> <output.docx>")
    build_document(sys.argv[1], sys.argv[2])
